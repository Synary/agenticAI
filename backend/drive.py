"""
Module de gestion de Google Drive
Synchronisation, extraction et indexation des fichiers Drive dans ChromaDB
"""
import json
import logging
import os
import pickle
import tempfile
import shutil
import hashlib
import io
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime

from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from langchain_core.tools import tool

from config import (
    logger, DRIVE_SCOPES, DRIVE_TOKEN_FILE, email_collection, knowledge_collection,
    TEXT_EXTENSIONS, AUDIO_EXTENSIONS, GOOGLE_DOC_MIME, GOOGLE_SHEET_MIME, GOOGLE_SLIDE_MIME,
    DEEPSEEK_API_KEY, processed_drive_files_file, client,model
)
from transcription import transcribe_audio_file


# --- AUTHENTIFICATION GOOGLE DRIVE ---
def get_drive_service():
    """Authentifie l'utilisateur et retourne le service Google Drive."""
    creds = None

    if os.path.exists(DRIVE_TOKEN_FILE):
        with open(DRIVE_TOKEN_FILE, "rb") as token:
            creds = pickle.load(token)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists("drive_credentials.json"):
                raise FileNotFoundError(
                    "drive_credentials.json not found. "
                    "Place it next to your backend."
                )

            flow = InstalledAppFlow.from_client_secrets_file(
                "drive_credentials.json",
                DRIVE_SCOPES
            )
            creds = flow.run_local_server(port=0)

        with open(DRIVE_TOKEN_FILE, "wb") as token:
            pickle.dump(creds, token)

    return build("drive", "v3", credentials=creds)


# --- CHARGEMENT/SAUVEGARDE DES FICHIERS DRIVE TRAITÉS ---
def load_processed_drive_files() -> dict:
    """Charge le suivi des fichiers Drive traités avec timestamp"""
    try:
        with open(processed_drive_files_file, 'r') as f:
            return json.load(f)
    except FileNotFoundError:
        with open(processed_drive_files_file, 'w') as f:
            json.dump({}, f)
        return {}
    except json.JSONDecodeError:
        with open(processed_drive_files_file, 'w') as f:
            json.dump({}, f)
        return {}


def save_processed_drive_files(processed_dict: dict):
    """Sauvegarde le suivi des fichiers Drive traités"""
    with open(processed_drive_files_file, 'w') as f:
        json.dump(processed_dict, f)


# --- UTILITAIRES DRIVE ---
def _list_drive_files(service, page_size: int = 100):
    """Fonction interne: liste tous les fichiers non supprimés de Google Drive."""
    files = []
    page_token = None

    while True:
        response = service.files().list(
            pageSize=page_size,
            pageToken=page_token,
            q="trashed = false",
            orderBy="modifiedTime desc",
            fields=(
                "nextPageToken,"
                "files("
                "id,name,mimeType,size,modifiedTime,webViewLink,parents"
                ")"
            )
        ).execute()

        files.extend(response.get("files", []))
        page_token = response.get("nextPageToken")

        if not page_token:
            break

    return files


def download_drive_file(service, file_id: str, destination: str):
    """Télécharge un fichier depuis Google Drive"""
    request = service.files().get_media(fileId=file_id)
    with open(destination, 'wb') as f:
        downloader = MediaIoBaseDownload(f, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()


def export_google_file(service, file_id: str, mime_type: str, destination: str):
    """Exporte un fichier Google Workspace (Docs, Sheets, Slides)"""
    export_mime = {
        "application/vnd.google-apps.document": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.google-apps.spreadsheet": "text/csv",
        "application/vnd.google-apps.presentation": "text/plain"
    }

    request = service.files().export_media(
        fileId=file_id,
        mimeType=export_mime.get(mime_type, "text/plain")
    )

    with open(destination, 'wb') as f:
        downloader = MediaIoBaseDownload(f, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()


def extract_text_from_pdf(file_path: str) -> str:
    """Extrait le texte d'un PDF"""
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text[:5000]
    except Exception as e:
        logger.warning(f"Could not extract PDF: {str(e)}")
        return ""


def extract_text_from_docx(file_path: str) -> str:
    """Extrait le texte d'un DOCX"""
    try:
        doc = Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text[:5000]
    except Exception as e:
        logger.warning(f"Could not extract DOCX: {str(e)}")
        return ""


def extract_text_from_xlsx(file_path: str) -> str:
    """Extrait le texte d'un XLSX"""
    try:
        workbook = load_workbook(file_path)
        content = ""
        for sheet in workbook.sheetnames[:2]:
            ws = workbook[sheet]
            for row in ws.iter_rows(values_only=True):
                content += str(row) + "\n"
        return content[:5000]
    except Exception as e:
        logger.warning(f"Could not extract XLSX: {str(e)}")
        return ""


# --- FONCTIONS D'EXTRACTION ET RÉSUMÉ ---
def get_file_category(filename: str, mime_type: str) -> str:
    """Détermine la catégorie d'un fichier"""
    ext = Path(filename).suffix.lower()

    if ext in TEXT_EXTENSIONS or mime_type.startswith("text/"):
        return "text"
    elif ext in AUDIO_EXTENSIONS or mime_type.startswith("audio/"):
        return "audio"
    elif mime_type in ("application/pdf",):
        return "pdf"
    elif mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",):
        return "docx"
    elif mime_type in ("text/csv", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
        return "spreadsheet"
    elif mime_type == GOOGLE_DOC_MIME:
        return "google_doc"
    elif mime_type == GOOGLE_SHEET_MIME:
        return "google_sheet"
    elif mime_type == GOOGLE_SLIDE_MIME:
        return "google_slide"
    else:
        return "unknown"


def extract_drive_content(file_path: str, filename: str, mime_type: str) -> tuple:
    """Extrait le contenu d'un fichier Drive"""
    category = get_file_category(filename, mime_type)
    content = ""

    try:
        if category == "text":
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()[:5000]
        elif category == "pdf":
            content = extract_text_from_pdf(file_path)
        elif category == "docx":
            content = extract_text_from_docx(file_path)
        elif category == "spreadsheet":
            content = extract_text_from_xlsx(file_path)
        elif category == "audio":
            with open(file_path, 'rb') as f:
                file_content = f.read()
            content = transcribe_audio_file(file_content, filename)
            content = content[:5000]
            category = "transcription"
        else:
            content = f"File: {filename} ({mime_type})"

    except Exception as e:
        logger.warning(f"Could not extract content from {filename}: {str(e)}")
        content = f"File: {filename}\n[Extraction error: {str(e)}]"

    return content, category


def save_drive_content_to_chroma(file_id: str, filename: str, mime_type: str,
                                 category: str, content: str, modified_time: str,
                                 web_view_link: str) -> int:
    """Sauvegarde le contenu d'un fichier Drive dans ChromaDB"""
    prefix = "🎙️" if category == "audio" else "📄"

    knowledge_collection.add(
        documents=[f"{prefix} {filename}\n\n{content}"],
        metadatas=[{
            "type": "drive_file",
            "filename": filename,
            "mime_type": mime_type,
            "category": category,
            "file_id": file_id,
            "modified_time": modified_time,
            "link": web_view_link
        }],
        ids=[f"drive_{file_id}"]
    )

    return 1  # 1 chunk


def remove_existing_drive_chunks(file_id: str):
    """Supprime les chunks existants d'un fichier Drive de ChromaDB"""
    try:
        knowledge_collection.delete(ids=[f"drive_{file_id}"])
        logger.info(f"Removed existing chunks for {file_id}")
    except Exception as e:
        logger.warning(f"Could not remove chunks for {file_id}: {str(e)}")


@tool
def list_drive_files(page_size: int = 100) -> str:
    """
    Liste les fichiers accessibles dans Google Drive.

    Args:
        page_size: Nombre maximum de fichiers récupérés par page.

    Returns:
        Liste formatée des fichiers Drive avec leurs métadonnées.
    """
    logger.info(f"📁 Tool list_drive_files called (page_size={page_size})")

    try:
        service = get_drive_service()
        files = _list_drive_files(service, page_size=page_size)

        if not files:
            return "Google Drive is accessible, but no untrashed files were found."

        formatted = []
        for index, drive_file in enumerate(files, 1):
            name = drive_file.get("name", "Untitled")
            file_id = drive_file.get("id", "")
            mime_type = drive_file.get("mimeType", "")
            modified_time = drive_file.get("modifiedTime", "")
            link = drive_file.get("webViewLink", "")

            formatted.append(
                f"{index}. **{name}**\n"
                f"   ID: {file_id}\n"
                f"   Type: {mime_type}\n"
                f"   Modified: {modified_time}\n"
                f"   Link: {link or 'not available'}"
            )

        return "\n\n".join(formatted)

    except Exception as e:
        logger.error(f"Error listing Drive files: {str(e)}")
        return f"Erreur: {str(e)}"


@tool
def summarize_drive_document(document_content: str, filename: str = "", file_type: str = "") -> str:
    """
    Résume le contenu d'un document Drive.

    Args:
        document_content: Le contenu textuel du document
        filename: Nom du fichier (optionnel)
        file_type: Type de fichier (optionnel)

    Returns:
        Un résumé concis du document
    """
    logger.info(f"Summarize tool called for Drive document: {filename}")

    content_preview = document_content[:5000]

    prompt = f"""Tu es un assistant qui résume des documents de manière concise et structurée.

    {f"FILE: {filename}" if filename else ""}
    {f"TYPE: {file_type}" if file_type else ""}

    CONTENT:
    {content_preview}

    RÉSUMÉ (3-5 phrases claires et structurées):
    """

    try:
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            max_tokens=500
        )
        return response.choices[0].message.content
    except Exception as e:
        logger.error(f"Error summarizing Drive document: {str(e)}")
        return f"Erreur lors du résumé: {str(e)}"


@tool
def sync_drive_tool(placeholder: str = "") -> str:
    """
    Synchronise Google Drive vers ChromaDB.
    Outil directement utilisable par l'agent.
    """
    logger.info("🚀 Agent tool: Google Drive sync requested")

    try:
        result = sync_google_drive()
        return (
            f"✅ Google Drive sync complete. "
            f"Total: {result['total_files']}, "
            f"Indexed: {result['indexed']}, "
            f"Already processed: {result['already_processed']}, "
            f"Skipped: {result['skipped']}, "
            f"Errors: {result['failed']}."
        )
    except Exception as e:
        logger.error(f"❌ Drive sync tool error: {e}", exc_info=True)
        return f"Error syncing Google Drive: {str(e)}"


def sync_google_drive():
    """
    Synchronise les fichiers supportés de Google Drive vers ChromaDB.
    """
    logger.info("🚀 Starting Google Drive synchronization...")

    processed_drive_files = load_processed_drive_files()

    service = get_drive_service()
    files = _list_drive_files(service)

    logger.info(f"📁 {len(files)} files found in Google Drive")

    indexed = 0
    skipped = 0
    failed = 0
    already_processed = 0
    results = []

    temp_dir = tempfile.mkdtemp(prefix="drive_sync_")

    try:
        for drive_file in files:
            file_id = drive_file.get("id", "")
            filename = drive_file.get("name", "")
            mime_type = drive_file.get("mimeType", "")
            modified_time = drive_file.get("modifiedTime", "")
            web_view_link = drive_file.get("webViewLink", "")

            logger.info(f"➡️ Drive: {filename} | {mime_type}")

            if file_id in processed_drive_files:
                stored_time = processed_drive_files[file_id].get("modified_time", "")
                if stored_time == modified_time:
                    logger.info(f"📌 {filename} already processed and unchanged, skipping...")
                    already_processed += 1
                    results.append({
                        "filename": filename,
                        "file_id": file_id,
                        "status": "already_processed"
                    })
                    continue
                else:
                    logger.info(f"🔄 {filename} was modified, reprocessing...")
                    remove_existing_drive_chunks(file_id)

            try:
                category = get_file_category(filename, mime_type)

                if category == "unknown":
                    skipped += 1
                    results.append({
                        "filename": filename,
                        "file_id": file_id,
                        "status": "skipped",
                        "reason": "unsupported_type"
                    })
                    continue

                is_google_workspace = mime_type.startswith("application/vnd.google-apps.")

                if mime_type == GOOGLE_DOC_MIME:
                    local_extension = ".docx"
                elif mime_type == GOOGLE_SHEET_MIME:
                    local_extension = ".csv"
                elif mime_type == GOOGLE_SLIDE_MIME:
                    local_extension = ".txt"
                else:
                    local_extension = Path(filename).suffix.lower()

                if not local_extension:
                    local_extension = ".bin"

                local_path = os.path.join(
                    temp_dir,
                    f"{hashlib.sha256(file_id.encode()).hexdigest()}{local_extension}"
                )

                if is_google_workspace:
                    export_google_file(service, file_id, mime_type, local_path)
                else:
                    download_drive_file(service, file_id, local_path)

                content, extracted_category = extract_drive_content(local_path, filename, mime_type)
                content = (content or "").strip()

                if not content:
                    skipped += 1
                    results.append({
                        "filename": filename,
                        "file_id": file_id,
                        "status": "skipped",
                        "reason": "empty_content"
                    })
                    continue

                chunks_count = save_drive_content_to_chroma(
                    file_id, filename, mime_type, extracted_category,
                    content, modified_time, web_view_link
                )

                indexed += 1
                processed_drive_files[file_id] = {
                    "filename": filename,
                    "modified_time": modified_time,
                    "processed_at": datetime.now().isoformat(),
                    "chunks_count": chunks_count
                }
                save_processed_drive_files(processed_drive_files)

                results.append({
                    "filename": filename,
                    "file_id": file_id,
                    "status": "indexed",
                    "chunks": chunks_count
                })

            except Exception as e:
                failed += 1
                logger.error(f"❌ Error with {filename}: {e}", exc_info=True)
                results.append({
                    "filename": filename,
                    "file_id": file_id,
                    "status": "error",
                    "error": str(e)
                })

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    logger.info(
        f"✅ Drive sync complete | "
        f"indexed={indexed}, already_processed={already_processed}, "
        f"skipped={skipped}, failed={failed}"
    )

    return {
        "total_files": len(files),
        "indexed": indexed,
        "already_processed": already_processed,
        "skipped": skipped,
        "failed": failed,
        "chromadb_total_chunks": knowledge_collection.count(),
        "results": results
    }


def fetch_and_store_recent_drive_files(max_files: int = 4) -> Dict[str, Any]:
    """
    Récupère les N derniers fichiers Drive et les stocke en ChromaDB

    Args:
        max_files: Nombre de fichiers à récupérer (défaut: 4)

    Returns:
        Dictionnaire avec stats et liste des fichiers traités
    """
    logger.info(f"🔄 Fetching and storing last {max_files} Drive files...")

    processed_drive_files = load_processed_drive_files()

    service = get_drive_service()
    files = _list_drive_files(service)[:max_files]

    files_stored = []
    temp_dir = tempfile.mkdtemp(prefix="drive_fetch_")

    try:
        for drive_file in files:
            file_id = drive_file.get("id", "")
            file_name = drive_file.get("name", "")
            mime_type = drive_file.get("mimeType", "")
            size = drive_file.get("size", 0)
            created_time = drive_file.get("createdTime", "")
            web_link = drive_file.get("webViewLink", "")

            if file_id in processed_drive_files:
                logger.info(f"⏭️ File {file_id} already processed")
                continue

            is_audio = False

            try:
                local_path = os.path.join(temp_dir, f"{file_id}_{file_name}")

                if mime_type.startswith("application/vnd.google-apps."):
                    if mime_type == GOOGLE_DOC_MIME:
                        export_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    elif mime_type == GOOGLE_SHEET_MIME:
                        export_mime = "text/csv"
                    else:
                        export_mime = "text/plain"

                    request = service.files().export_media(fileId=file_id, mimeType=export_mime)
                else:
                    request = service.files().get_media(fileId=file_id)

                with open(local_path, 'wb') as f:
                    downloader = MediaIoBaseDownload(f, request)
                    done = False
                    while not done:
                        status, done = downloader.next_chunk()

                content = ""
                if mime_type == "application/pdf":
                    content = extract_text_from_pdf(local_path)
                elif mime_type.endswith(".docx") or mime_type == GOOGLE_DOC_MIME:
                    content = extract_text_from_docx(local_path)
                elif "spreadsheet" in mime_type.lower() or mime_type == GOOGLE_SHEET_MIME:
                    content = extract_text_from_xlsx(local_path)
                elif mime_type.startswith("audio/"):
                    is_audio = True
                    with open(local_path, 'rb') as f:
                        file_content = f.read()
                    content = transcribe_audio_file(file_content, file_name)
                    content = content[:5000]
                else:
                    content = f"File: {file_name} ({mime_type}) - {size} bytes"

            except Exception as e:
                logger.warning(f"Could not extract content from {file_name}: {str(e)}")
                content = f"File: {file_name}\n[Error: {str(e)}]"

            drive_id = f"drive_{file_id}"
            metadata = {
                "type": "drive_file_audio" if is_audio else "drive_file",
                "filename": file_name,
                "mime_type": mime_type,
                "size": size,
                "created_at": created_time,
                "file_id": file_id,
                "link": web_link,
                "is_transcribed": is_audio
            }

            prefix = "🎙️" if is_audio else "📄"
            email_collection.add(
                documents=[f"{prefix} {file_name}\n\n{content}"],
                metadatas=[metadata],
                ids=[drive_id]
            )

            logger.info(f"✅ Drive file stored: {drive_id}")

            files_stored.append({
                "id": file_id,
                "name": file_name,
                "mime_type": mime_type,
                "size": size,
                "link": web_link,
                "is_audio": is_audio
            })

            processed_drive_files[file_id] = {
                "name": file_name,
                "mime_type": mime_type,
                "is_audio": is_audio,
                "processed_at": datetime.now().isoformat()
            }

        save_processed_drive_files(processed_drive_files)

        return {
            "status": "success",
            "count": len(files_stored),
            "files": files_stored,
            "chromadb_total": email_collection.count()
        }

    except Exception as e:
        logger.error(f"❌ Error fetching Drive files: {str(e)}")
        return {
            "status": "error",
            "error": str(e)
        }

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
